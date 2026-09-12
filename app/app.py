
# app.py - MVP Streamlit app for Diagnostic EFA/OPA
import streamlit as st
import uuid, json, sqlite3, os
from datetime import date
from io import BytesIO

try:
    from docx import Document
except Exception:
    Document = None

import pandas as pd

BASE_DIR = "data"
AUDIO_DIR = os.path.join(BASE_DIR, "audio")
DB_PATH = os.path.join(BASE_DIR, "diagnostics.db")
os.makedirs(AUDIO_DIR, exist_ok=True)
os.makedirs(BASE_DIR, exist_ok=True)

# --- DB init (SQLite) ---
def init_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS diagnostics (
                 id TEXT PRIMARY KEY, payload TEXT, created DATE)""")
    conn.commit()
    return conn

conn = init_db()

def save_diagnostic(payload: dict):
    id_ = str(uuid.uuid4())
    cur = conn.cursor()
    cur.execute("INSERT INTO diagnostics (id,payload,created) VALUES (?,?,?)",
                (id_, json.dumps(payload, ensure_ascii=False), str(date.today())))
    conn.commit()
    return id_

# --- Stub LLM caller (remplacer par Ollama/OpenAI) ---
def call_llm(prompt:str, model="local-stub"):
    # Retourner une structure JSON: {"summary": "...", "analysis": {...}}
    # Remplacez cette fonction par un appel à Ollama/OpenAI.
    return {
        "summary": "Résumé automatique (stub) — remplacez call_llm par un vrai appel LLM.",
        "analysis": {
            "SWOT": {
                "Forces": ["Force A (exemple)"],
                "Faiblesses": ["Faiblesse A (exemple)"],
                "Opportunites": ["Opportunité A (exemple)"],
                "Menaces": ["Menace A (exemple)"]
            },
            "PESTEL": {"P": [], "E": [], "S": [], "T": [], "E2": [], "L": []},
            "Porter": {"Entrants": [], "Clients": [], "Substituts": [], "Rivalite": [], "Fournisseurs": []},
            "BCG": [],
            "Ansoff": []
        }
    }

def generate_docx_report(payload: dict, analysis: dict) -> bytes:
    """Génère un rapport DOCX en mémoire et retourne les bytes (ou None si python-docx absent)."""
    if Document is None:
        return None
    doc = Document()
    doc.add_heading("Plan stratégique - Diagnostic EFA/OPA", level=1)
    doc.add_paragraph(f"Nom: {payload.get('nom', '')}")
    doc.add_paragraph(f"Type: {payload.get('type', '')}   Date: {payload.get('date', '')}")
    doc.add_heading("Synthèse LLM", level=2)
    doc.add_paragraph(analysis.get("summary", ""))
    doc.add_heading("Analyse détaillée", level=2)
    doc.add_paragraph("SWOT")
    sw = analysis.get("analysis", {}).get("SWOT", {})
    for k in ["Forces","Faiblesses","Opportunites","Menaces"]:
        doc.add_heading(k, level=3)
        items = sw.get(k, [])
        for it in items:
            doc.add_paragraph(f"- {it}")
    # autres sections (PESTEL, Porter...)
    doc.add_page_break()
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()

st.set_page_config(page_title="Agent IA - Diagnostic EFA/OPA", layout="wide")
st.title("Agent IA — Diagnostic & Analyse stratégique (EFA / OPA)")

with st.form("diagnostic_form"):
    col1, col2 = st.columns([2,1])
    with col1:
        type_diag = st.selectbox("Type de diagnostic", ["EFA","OPA"])
        nom = st.text_input("Nom de l'exploitation / OPA")
        responsable = st.text_input("Responsable / contact")
        date_diag = st.date_input("Date du diagnostic", value=date.today())
        st.header("Branches - étoile du conseil")
        moyens = st.text_area("Moyens de production (description)")
        perf = st.text_area("Performances technico-économiques (résumé par activité)")
        finances_text = st.text_area("Finances (chiffres clés, remarques)")
        milieu = st.text_area("Milieu local")
        marches = st.text_area("Marchés / filières")
        politiques = st.text_area("Politiques publiques")
    with col2:
        st.header("Pièces & enregistrements")
        audio_file = st.file_uploader("Uploader un enregistrement audio (wav/mp3)", type=["wav","mp3"])
        saved_audio_path = None
        if audio_file is not None:
            fname = f"{uuid.uuid4()}_{audio_file.name}"
            path = os.path.join(AUDIO_DIR, fname)
            with open(path, "wb") as f:
                f.write(audio_file.read())
            saved_audio_path = path
            st.success("Audio sauvegardé.")
        st.header("Options")
        validation = st.checkbox("Validation préalable par le conseiller (cocher si ok)", value=False)

    submitted = st.form_submit_button("Enregistrer diagnostic")
if submitted:
    payload = {
        "type": type_diag,
        "nom": nom,
        "responsable": responsable,
        "date": date_diag.isoformat(),
        "branches": {
            "moyens_production": moyens,
            "performances_techno_economiques": perf,
            "finances": finances_text,
            "milieu_local": milieu,
            "marches_filieres": marches,
            "politiques_publiques": politiques
        },
        "validation_conseiller": validation,
        "audio_file": saved_audio_path
    }
    id_ = save_diagnostic(payload)
    st.success(f"Diagnostic enregistré (id={id_}).")
    st.info("Appuyez sur 'Analyser' pour lancer l'analyse stratégique sur ce diagnostic.")

st.markdown("---")
st.header("Analyse stratégique")
diag_id = st.text_input("Entrez l'ID du diagnostic à analyser (ou vide pour analyser le dernier)", value="")
if st.button("Analyser"):
    cur = conn.cursor()
    if not diag_id:
        cur.execute("SELECT id,payload FROM diagnostics ORDER BY created DESC LIMIT 1")
    else:
        cur.execute("SELECT id,payload FROM diagnostics WHERE id=? LIMIT 1", (diag_id,))
    row = cur.fetchone()
    if not row:
        st.error("Diagnostic introuvable.")
    else:
        payload = json.loads(row[1])
        prompt = f"Analyse stratégique (FR) pour le diagnostic suivant:\\n{json.dumps(payload, ensure_ascii=False)}\\nProduis: SWOT, PESTEL, Porter(5forces), BCG, Ansoff, actions priorisées et plan stratégique 3-5 ans et plan d'action annuel."
        llm_out = call_llm(prompt)
        st.subheader("Synthèse LLM")
        st.write(llm_out.get("summary"))
        st.subheader("Analyse détaillée")
        st.json(llm_out.get("analysis"))
        # Proposer export DOCX si possible
        report_bytes = generate_docx_report(payload, llm_out) if Document is not None else None
        if report_bytes:
            st.download_button("Télécharger rapport DOCX", data=report_bytes, file_name=f"plan_strategique_{payload.get('nom','')}.docx")
        else:
            st.info("python-docx non disponible dans l'environnement serveur -> export DOCX non proposé.")
